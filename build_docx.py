#!/usr/bin/env python3
"""Generate a .docx documentation file for the university-collab-network project."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

import os
import re

ROOT = os.path.dirname(os.path.abspath(__file__))

ACCENT = RGBColor(0x1B, 0x5E, 0x20)
DARK = RGBColor(0x21, 0x21, 0x21)
GRAY = RGBColor(0x55, 0x55, 0x55)

APP = os.path.join(ROOT, "apps")
API = os.path.join(APP, "api")
WEB = os.path.join(APP, "web")
SHARED = os.path.join(ROOT, "packages", "shared-types")


def read(rel):
    p = os.path.join(ROOT, rel)
    try:
        with open(p, "r", encoding="utf-8", errors="replace") as f:
            return f.read()
    except OSError:
        return ""


def models_and_enums():
    schema = read("apps/api/prisma/schema.prisma")
    models = re.findall(r"^model\s+(\w+)", schema, re.M)
    enums = re.findall(r"^enum\s+(\w+)", schema, re.M)
    return models, enums


def route_table():
    rows = []
    routes_dir = os.path.join(API, "src", "routes")
    for fname in sorted(os.listdir(routes_dir)):
        if not fname.endswith(".routes.ts"):
            continue
        text = read(os.path.join("apps/api/src/routes", fname))
        base = fname.replace(".routes.ts", "")
        entries = re.findall(r"\.(get|post|put|patch|delete)\s*\(\s*[\"']([^\"']*)[\"']", text, re.I)
        seen = set()
        for method, sub in entries:
            key = (method.upper(), sub or "/")
            if key in seen:
                continue
            seen.add(key)
            rows.append((f"/api/v1/{base}", method.upper(), f"`{sub or '/'}`"))
    return rows


def register_fonts(doc):
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(4)


def h1(doc, text):
    p = doc.add_heading(level=1)
    r = p.add_run(text)
    r.font.color.rgb = ACCENT
    r.font.size = Pt(20)
    return p


def h2(doc, text):
    p = doc.add_heading(level=2)
    r = p.add_run(text)
    r.font.color.rgb = ACCENT
    r.font.size = Pt(15)
    return p


def h3(doc, text):
    p = doc.add_heading(level=3)
    r = p.add_run(text)
    r.font.color.rgb = DARK
    r.font.size = Pt(12)
    return p


def para(doc, text, italic=False, bold=False, color=None, size=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = italic
    r.bold = bold
    if color:
        r.font.color.rgb = color
    if size:
        r.font.size = Pt(size)
    return p


def bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.add_run(text)
    return p


def code_block(doc, text):
    for line in text.splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.4)
        r = p.add_run(line)
        r.font.name = "Consolas"
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x1F, 0x1F, 0x1F)
        p.paragraph_format.space_after = Pt(0)


def main():
    doc = Document()
    register_fonts(doc)

    # Title page
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("University Collaboration & Research Network")
    r.bold = True
    r.font.size = Pt(28)
    r.font.color.rgb = ACCENT

    st = doc.add_paragraph()
    st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = st.add_run("Full Codebase Walkthrough & Technical Documentation")
    r.font.size = Pt(14)
    r.font.color.rgb = GRAY

    doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(
        "Monorepo: pnpm workspaces + Turborepo  •  "
        "Backend: Express + TypeScript + Prisma + Postgres  •  "
        "Frontend: React 18 + Vite + Tailwind + Zustand + TanStack Query"
    ).font.size = Pt(10)

    doc.add_page_break()

    # 1. Overview
    h1(doc, "1. What Is This Project?")
    para(
        doc,
        "A university-scoped platform connecting Students, Professors, Researchers, "
        "Research Teams, Projects, Clubs/Societies, Startups, Publications, Research "
        "Topics, Opportunities, and Events into one discoverable collaboration graph.",
    )
    para(
        doc,
        "Core question it answers for every user: \u201cWho / what project / what research / "
        "what opportunity should I know about right now?\u201d",
        italic=True,
    )
    para(
        doc,
        "It is NOT an ERP and NOT a social-media clone. Every feature must trace back to: "
        "discovery, collaboration, research, networking, opportunity-matching, or security "
        "(the \u201cAnti-Crap Rule\u201d in PROJECT_SPEC.md).",
    )

    # 2. Stack
    h1(doc, "2. Technology Stack")
    table = doc.add_table(rows=1, cols=2)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text = "Layer", "Technology"
    stack = [
        ("Monorepo", "pnpm workspaces + Turborepo"),
        ("Backend API", "Node 20, Express 4, TypeScript"),
        ("Database", "PostgreSQL via Prisma ORM 5"),
        ("Validation", "Zod (schemas shared web + api)"),
        ("Auth", "bcrypt + JWT (access 15m / refresh 7d, httpOnly cookie)"),
        ("Frontend", "React 18, Vite 5, TypeScript"),
        ("UI/Styling", "Tailwind CSS 3, CVA, clsx, tailwind-merge"),
        ("Server state", "TanStack Query 5 (hooks per resource)"),
        ("Client state", "Zustand 5 (session, ui, toast stores)"),
        ("Routing", "react-router-dom 6"),
        ("Logging", "pino + pino-http"),
        ("Tooling", "ESLint 9, Prettier 3, Husky, TypeScript 5.6"),
    ]
    for row in stack:
        cells = table.add_row().cells
        cells[0].text = row[0]
        cells[1].text = row[1]

    # 3. Folder structure
    h1(doc, "3. Folder Structure")
    code_block(doc, read("FILE_STRUCTURE.md").split("Monorepo Tree")[0])
    code_block(doc, read("FILE_STRUCTURE.md").split("Monorepo Tree")[1].split("Route \u2192")[0].strip()[:2200])

    para(doc, "")

    # 4. Architecture / layering
    h1(doc, "4. Architecture & Data Flow")
    h2(doc, "Layered architecture")
    para(
        doc,
        "UI (pages) \u2192 Components \u2192 Hooks/State (TanStack Query + Zustand) \u2192 "
        "Services/API client \u2192 Backend (Routes \u2192 Controllers \u2192 Services \u2192 "
        "Repositories (Prisma) \u2192 Database).",
    )
    bullet(doc, "Business logic lives ONLY in backend services/*")
    bullet(doc, "Controllers only parse/validate (Zod) and delegate to a service")
    bullet(doc, "Repositories isolate every Prisma query")
    bullet(doc, "Frontend never talks to Prisma and contains no authorization logic")

    h2(doc, "A request, end to end")
    for step in [
        "1. A React component calls a typed client in apps/web/src/services/api/ (never ad-hoc fetch).",
        "2. Client sends JSON to http://localhost:4000/api/v1/* with Authorization: Bearer <access_jwt>.",
        "3. Express middleware chain (apps/api/src/app.ts): logging \u2192 CORS \u2192 body parse \u2192 cookies \u2192 /health \u2192 /api/v1 router \u2192 404 \u2192 errorHandler.",
        "4. Versioned router (routes/index.ts) dispatches to the matching resource router.",
        "5. Controller validates with a Zod schema, then calls a service.",
        "6. Service applies business rules (privacy filtering, match scoring) and calls a repository.",
        "7. Repository runs the Prisma query against Postgres and returns data.",
        "8. JSON response returns up the same chain.",
    ]:
        para(doc, step)

    # 5. Auth flow
    h1(doc, "5. Authentication & Authorization")
    for line in [
        "Signup \u2192 users row created with status=pending_verification.",
        "Email verification (university domains auto-verified via UNIVERSITY_EMAIL_DOMAINS).",
        "Login \u2192 bcrypt check \u2192 JWT access token (15m) + refresh token (7d, httpOnly cookie).",
        "Refresh \u2192 rotates token, old one revoked (stored hashed in refresh_tokens).",
        "Middleware: requireAuth (validates JWT, sets req.user); requireRole (checks verifications table, never requested_role alone).",
        "Privacy: public | university_only | connections_only | private, with independent field-level controls. Private fields are OMITTED from API responses at the service layer (DECISIONS D-004) \u2014 the frontend is never trusted.",
    ]:
        bullet(doc, line)

    # 6. Models
    models, enums = models_and_enums()
    h1(doc, "6. Data Model (Prisma)")
    para(doc, f"{len(models)} models | {len(enums)} enums | ~911-line schema.prisma")
    h3(doc, "Models")
    cols = doc.add_table(rows=1, cols=1)
    cols.style = "Light List Accent 1"
    cell = cols.rows[0].cells[0]
    for i, name in enumerate(models, 1):
        cell.text += f"{i}. {name}\n" if i == 1 else f"{i}. {name}\n"
    cell.text = cell.text.strip()
    h3(doc, "Core relationship web (PROJECT_SPEC.md \u00a75)")
    code_block(doc, """Student  -MEMBER_OF->       Organization  (club/society/startup)
Professor -LEADS->             ResearchTeam
Researcher-MEMBER_OF->         ResearchTeam
Researcher-AUTHORS->           Publication
Publication-ABOUT->            ResearchTopic
Project  -RELATED_TO->         ResearchTopic
Project  -USES->               Skill
Opportunity-PROVIDED_BY->     Organization | ResearchTeam
Event    -ORGANIZED_BY->       Organization | ResearchTeam""")

    # 7. API routes
    h1(doc, "7. API Surface (mounted under /api/v1)")
    rows = route_table()
    tbl = doc.add_table(rows=1, cols=3)
    tbl.style = "Light Grid Accent 1"
    hdr = tbl.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = "Base path", "Method", "Path"
    for row in rows:
        cells = tbl.add_row().cells
        cells[0].text = row[0]
        cells[1].text = row[1]
        cells[2].text = row[2]
    para(doc, "")
    para(doc, "Shared conventions (API_CONTRACT.md):")
    bullet(doc, "Errors: { \"error\": { code, message, fields? } }")
    bullet(doc, "Pagination: ?cursor=&limit= (default 20, max 100) \u2192 { data, nextCursor }")
    bullet(doc, "Auth header: Authorization: Bearer <access_token> on all non-public routes")

    # 8. Matching algo
    h1(doc, "8. Team-Builder Match Score")
    para(doc, "services/matching.service.ts \u2014 weighted, explainable scoring:", )
    code_block(doc, """score = w1*skill_overlap          (0.40)  project skills matched
     + w2*research_topic_overlap  (0.30)
     + w3*availability_match      (0.15)
     + w4*department_proximity    (0.10)
     + w5*existing_connection     (0.05)""")
    para(
        doc,
        "Every match response includes the matched/unmatched criteria list (e.g. "
        "\u2713 Python, \u2713 ML, \u2717 React) next to the numeric score \u2014 a bare "
        "percentage without explanation is a spec violation.",
    )

    # 9. Status
    h1(doc, "9. Implementation Status")
    table = doc.add_table(rows=1, cols=2)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text = "Phase", "Status"
    text = read("IMPLEMENTATION_STATUS.md")
    for line in text.splitlines():
        m = re.match(r"### (Phase [\d.]+[^—]*)", line)
        s = re.search(r"\bSTATUS:\s*\*?\*?([^*\n]+)", line)
        if m and s:
            cells = table.add_row().cells
            cells[0].text = m.group(1).strip()
            cells[1].text = s.group(1).strip().split("(")[0].strip()[:120]
    para(doc, "")
    para(
        doc,
        "Note: some later phases were built ahead of the strict phase-gate order at the "
        "user's direction. Live DB / browser verification is still pending a working "
        "PostgreSQL. See AGENT_HANDOFF.md and DECISIONS.md for full history.",
        italic=True,
        color=GRAY,
    )

    # 10. How to run
    h1(doc, "10. How to Run & Develop")
    code_block(doc, """pnpm install
pnpm dev                      # run api + web together (turbo)

# backend
cd apps/api
  pnpm prisma:generate         # regenerate Prisma client
  pnpm prisma:migrate          # apply migrations to Postgres
  pnpm prisma:seed             # interconnected demo data
  pnpm prisma:studio           # DB GUI
  pnpm test                    # backend tests

# frontend
cd apps/web
  pnpm dev                     # Vite dev server (port 5173)
  pnpm build                   # production bundle""")
    h3(doc, "Required env vars (apps/api/.env)")
    env = read("apps/api/.env.example")
    names = re.findall(r"(?m)^([A-Z0-9_]+)=", env)
    para(doc, ", ".join(names))
    para(doc, "You need a running PostgreSQL; set DATABASE_URL in apps/api/.env.")

    doc.add_page_break()
    h1(doc, "Appendix: Key Documentation Files")
    for f in ["PROJECT_SPEC.md", "ARCHITECTURE.md", "API_CONTRACT.md", "DATABASE_SCHEMA.md", "DEPENDENCIES.md", "ENVIRONMENT.md", "DECISIONS.md", "AGENT_HANDOFF.md", "IMPLEMENTATION_STATUS.md", "TESTING.md", "FILE_STRUCTURE.md"]:
        bullet(doc, f, level=0)

    out = os.path.join(ROOT, "UNIVERSITY_COLLAB_WALKTHROUGH.docx")
    doc.save(out)
    print("Saved:", out)


if __name__ == "__main__":
    main()